"""高校ごとの副申書の要否と基準をまとめる「副申書 必要条件一覧」Excel を生成する"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
data = json.loads((ROOT / "data" / "fukushinsho_criteria.json").read_text(encoding="utf-8"))

FONT = "Yu Gothic"
NAVY = "1F3A66"
FIXED = PatternFill("solid", fgColor="EEF1F4")    # 高校側で入力しない列
HONKOU = PatternFill("solid", fgColor="E3ECF7")   # 本校基準の行
WHITE = PatternFill("solid", fgColor="FFFFFF")
side = Side(style="thin", color="C3CBD4")
BORDER = Border(top=side, bottom=side, left=side, right=side)
HEAD = dict(font=Font(name=FONT, bold=True, color="FFFFFF", size=10),
            fill=PatternFill("solid", fgColor=NAVY),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
BODY_FONT = Font(name=FONT, size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(vertical="center", wrap_text=True)


def style(cell, fill=WHITE, font=BODY_FONT, alignment=LEFT):
    cell.border = BORDER
    cell.fill = fill
    cell.font = font
    cell.alignment = alignment


wb = Workbook()

# ── 使い方 ─────────────────────────────
us = wb.active
us.title = "使い方"
us.column_dimensions["A"].width = 2
us.column_dimensions["B"].width = 18
us.column_dimensions["C"].width = 100
us["B2"] = data["title"]
us["B2"].font = Font(name=FONT, bold=True, size=14)
us["B3"] = ("高校ごとに、欠席日数が多い生徒などの「副申書」が必要かどうかと、その基準をまとめる一覧です。"
            "各高校の担当の先生が、高校に確認して入力します。")
us["B3"].font = Font(name=FONT, bold=True, size=11)
us["B3"].alignment = Alignment(vertical="top", wrap_text=True)
us.merge_cells("B3:C3")
us.row_dimensions[3].height = 32
howto = [
    ("本校基準", "一覧の先頭行「戸町中（本校基準）」は、本校が副申書をつける生徒の基準です。"
             "各学年で欠席20日以上、成績に（　）表記や評定不能がある、試験で配慮が必要、のいずれかに当てはまる生徒が対象です。"),
    ("入力する人", "各高校の担当（「担当」列）。わからない学校は進路担当へ。"),
    ("入力するセル", "「必要条件一覧」の白いセル。要否は▼から「必要／不要／条件による／未確認」を選ぶ。"),
    ("欠席日数の基準", "「３年間の欠席日数」「各学年での欠席日数」「３年時の欠席日数」を、それぞれ「○日以上」の数字で入力。"
                "基準のない欄は空欄。いずれかに当てはまれば対象とみなします。日数にかかわらず必要なら空欄のまま「必要」。"),
    ("評定不能", "成績に（　）表記や評定不能がある生徒に副申書が必要なら▼から「対象」、不要なら「対象外」を選ぶ。"),
    ("健康上の配慮", "健康上の理由などで試験に配慮が必要な生徒に副申書が必要なら「対象」、不要なら「対象外」を選ぶ。"),
    ("条件による", "「理由が病気の場合のみ」など、上の列で表せない条件があるときに選び、「条件の補足」に書く。"),
    ("公立高校", "公立は県内共通の扱いなので、「公立高校（全校共通）」の1行にまとめます。"),
    ("一覧にない高校", "下の空いている行に高校名（共通マスタの表記）を入れて追加。"),
    ("統合アプリ", "このファイルを進路事務 統合アプリに読み込むと、決定届の志望校と成績データの欠席日数から"
              "「副申書が必要な生徒」を一覧にします（決定届タブ → 副申書が必要な生徒）。"),
]
for i, (k, v) in enumerate(howto, 4):
    us.cell(i, 2, k).font = Font(name=FONT, bold=True, size=11)
    us.cell(i, 2).alignment = Alignment(vertical="top")
    us.cell(i, 3, v).font = Font(name=FONT, size=11)
    us.cell(i, 3).alignment = Alignment(vertical="top", wrap_text=True)

# ── 必要条件一覧 ─────────────────────────────
ws = wb.create_sheet("必要条件一覧")
ws["A1"] = data["title"]
ws["A1"].font = Font(name=FONT, bold=True, size=14)
ws["A2"] = "白いセル＝入力。要否・評定不能・健康上の配慮が「未確認」の高校は、統合アプリで「確認が必要な高校」として表示されます。"
ws["A2"].font = Font(name=FONT, size=10, color="4A5868")

headers = [
    ("区分", 7), ("高校名", 24), ("担当", 9), ("副申書の要否", 12),
    ("３年間の欠席日数\n（以上）", 15), ("各学年での欠席日数\n（以上）", 15), ("３年時の欠席日数\n（以上）", 15),
    ("評定不能\n（（　）表記を含む）", 15), ("健康上の配慮\n（試験での配慮）", 15),
    ("条件の補足", 30), ("提出方法・様式", 26), ("確認日", 9), ("確認者", 9),
]
HR = 4
for c, (h, w) in enumerate(headers, 1):
    style(ws.cell(HR, c, h), **HEAD)
    ws.column_dimensions[ws.cell(HR, c).column_letter].width = w
ws.row_dimensions[HR].height = 44

# 本校基準
hk = data["honkou"]
r = HR + 1
row = ["本校", hk["name"], "", "", None, hk["each_grade_days"], None,
       hk["hyotei_funou"], hk["kenkou"], hk["note"], "", "", ""]
for c, v in enumerate(row, 1):
    style(ws.cell(r, c, v), fill=HONKOU, font=Font(name=FONT, size=10, bold=c <= 2),
          alignment=CENTER if c in (1, 5, 6, 7, 8, 9) else LEFT)
ws.row_dimensions[r].height = 58

# 高校ごとの行（空行を含めて30行目まで）
first, last = HR + 2, 30
for i, r in enumerate(range(first, last + 1)):
    kubun, name, tanto = data["schools"][i] if i < len(data["schools"]) else ("", "", "")
    filled = bool(name)
    for c, v in enumerate([kubun, name, tanto], 1):
        style(ws.cell(r, c, v or None), fill=FIXED if filled else WHITE,
              alignment=CENTER if c != 2 else LEFT)
    for c in range(4, 14):
        v = "未確認" if filled and c in (4, 8, 9) else None
        style(ws.cell(r, c, v), alignment=CENTER if c in (4, 5, 6, 7, 8, 9, 12, 13) else LEFT)
    ws.cell(r, 12).number_format = "m/d"

yohi = DataValidation(type="list", formula1='"必要,不要,条件による,未確認"', allow_blank=True)
taisho = DataValidation(type="list", formula1='"対象,対象外,未確認"', allow_blank=True)
days = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                      error="0以上の整数（日数）を入力してください。", showErrorMessage=True)
date = DataValidation(type="date", operator="greaterThanOrEqual", formula1="45000", allow_blank=True)
for dv, rng in [(yohi, f"D{first}:D{last}"), (taisho, f"H{first}:I{last}"),
                (days, f"E{first}:G{last}"), (date, f"L{first}:L{last}")]:
    ws.add_data_validation(dv)
    dv.add(rng)

AMBER = dict(fill=PatternFill("solid", bgColor="FFF2CC"), font=Font(color="9A6A0C"))
RED = dict(fill=PatternFill("solid", bgColor="FBEBED"), font=Font(color="A8323F"))
for rng, top, hit in [(f"D{first}:D{last}", "D", "必要"), (f"H{first}:I{last}", "H", "対象")]:
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{top}{first}="未確認"'], **AMBER))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{top}{first}="{hit}"'], **RED))

ws.freeze_panes = f"C{first}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

out = ROOT / "output" / "副申書_必要条件一覧.xlsx"
wb.active = 1
wb.save(out)
print(out)
