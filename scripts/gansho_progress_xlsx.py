"""私立高校 Web出願の進捗をクラス一覧で管理する Excel を生成する"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
data = json.loads((ROOT / "data" / "checklist.json").read_text(encoding="utf-8"))

STUDENTS = 40
FONT = "Yu Gothic"
NAVY = "1F4E79"
thin = Side(style="thin", color="808080")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
HEAD = dict(font=Font(name=FONT, bold=True, color="FFFFFF", size=10),
            fill=PatternFill("solid", fgColor=NAVY),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
BODY_FONT = Font(name=FONT, size=10)


def style(cell, **kw):
    cell.border = BORDER
    cell.font = kw.get("font", BODY_FONT)
    if "fill" in kw:
        cell.fill = kw["fill"]
    cell.alignment = kw.get("alignment", Alignment(vertical="center", wrap_text=True))


wb = Workbook()

# ── 進捗一覧 ─────────────────────────────
ws = wb.active
ws.title = "進捗一覧"
ws["A1"] = data["title"] + "　進捗一覧"
ws["A1"].font = Font(name=FONT, bold=True, size=14)
ws["A2"] = "各欄で「済・未・要確認」を選びます。「要確認」は赤、「済」は緑で表示されます。"
ws["A2"].font = Font(name=FONT, size=9, color="555555")

fixed = ["番号", "氏名", "受験校", "入試区分"]
ms = data["milestones"]
headers = fixed + ms + ["完了数", "進み具合"]
HR = 4
for c, h in enumerate(headers, 1):
    style(ws.cell(HR, c, h), **HEAD)
ws.row_dimensions[HR].height = 42

first_ms, last_ms = len(fixed) + 1, len(fixed) + len(ms)
fm, lm = get_column_letter(first_ms), get_column_letter(last_ms)
done_col, pct_col = last_ms + 1, last_ms + 2
first_row, last_row = HR + 1, HR + STUDENTS

center = Alignment(horizontal="center", vertical="center")
for i in range(STUDENTS):
    r = first_row + i
    style(ws.cell(r, 1, i + 1), alignment=center)
    for c in range(2, last_ms + 1):
        style(ws.cell(r, c), alignment=center if c >= first_ms else Alignment(vertical="center"))
    style(ws.cell(r, done_col, f'=IF(B{r}="","",COUNTIF({fm}{r}:{lm}{r},"済"))'), alignment=center)
    pc = ws.cell(r, pct_col, f"=IF(B{r}=\"\",\"\",{get_column_letter(done_col)}{r}/{len(ms)})")
    style(pc, alignment=center)
    pc.number_format = "0%"

# 集計行
sr = last_row + 1
style(ws.cell(sr, 1, "済の人数"), font=Font(name=FONT, bold=True, size=10),
      fill=PatternFill("solid", fgColor="DDEBF7"))
ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=len(fixed))
for c in range(first_ms, last_ms + 1):
    L = get_column_letter(c)
    style(ws.cell(sr, c, f'=COUNTIF({L}{first_row}:{L}{last_row},"済")'),
          fill=PatternFill("solid", fgColor="DDEBF7"), alignment=center)

dv = DataValidation(type="list", formula1='"済,未,要確認"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"{fm}{first_row}:{lm}{last_row}")
kubun = DataValidation(type="list", formula1='"推薦,専願,一般"', allow_blank=True)
ws.add_data_validation(kubun)
kubun.add(f"D{first_row}:D{last_row}")

rng = f"{fm}{first_row}:{lm}{last_row}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"済"'],
                              fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(color="006100")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"要確認"'],
                              fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006")))
pct_rng = f"{get_column_letter(pct_col)}{first_row}:{get_column_letter(pct_col)}{last_row}"
ws.conditional_formatting.add(pct_rng, FormulaRule(
    formula=[f'AND(ISNUMBER({get_column_letter(pct_col)}{first_row}),{get_column_letter(pct_col)}{first_row}=1)'],
    fill=PatternFill("solid", fgColor="C6EFCE")))

widths = [5, 14, 18, 9] + [8.5] * len(ms) + [7, 8]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = ws.cell(first_row, first_ms)
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = f"{HR}:{HR}"

# ── 受験校日程 ─────────────────────────────
ws2 = wb.create_sheet("受験校日程")
ws2["A1"] = "受験校ごとの日程（募集要項から転記）"
ws2["A1"].font = Font(name=FONT, bold=True, size=14)
cols2 = data["schoolColumns"] + ["事前入力の開始日", "中学校の校内締切", "受験者数", "メモ"]
for c, h in enumerate(cols2, 1):
    style(ws2.cell(3, c, h), **HEAD)
for r in range(4, 24):
    for c in range(1, len(cols2) + 1):
        style(ws2.cell(r, c))
    ws2.cell(r, len(cols2) - 1).value = f'=IF(A{r}="","",COUNTIF(進捗一覧!$C${first_row}:$C${last_row},A{r}))'
    ws2.cell(r, len(cols2) - 1).alignment = center
for c, w in enumerate([18, 14, 10, 14, 14, 11, 11, 14, 14, 14, 9, 24], 1):
    ws2.column_dimensions[get_column_letter(c)].width = w
ws2.row_dimensions[3].height = 32
ws2.freeze_panes = "B4"

# ── チェック項目（参照用） ─────────────────────────────
ws3 = wb.create_sheet("チェック項目")
ws3["A1"] = "チェック項目一覧（生徒配布用シートと同じ内容）"
ws3["A1"].font = Font(name=FONT, bold=True, size=14)
for c, h in enumerate(["No", "確認項目", "補足・メモ"], 1):
    style(ws3.cell(3, c, h), **HEAD)
r, n = 4, 1
for sec in data["sections"]:
    ws3.cell(r, 1, sec["title"])
    ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for c in range(1, 4):
        style(ws3.cell(r, c), font=Font(name=FONT, bold=True, size=10), fill=PatternFill("solid", fgColor="DDEBF7"))
    r += 1
    for item, note in sec["items"]:
        style(ws3.cell(r, 1, n), alignment=center)
        style(ws3.cell(r, 2, item))
        style(ws3.cell(r, 3, note), font=Font(name=FONT, size=9, color="555555"))
        r, n = r + 1, n + 1
for c, w in enumerate([5, 64, 34], 1):
    ws3.column_dimensions[get_column_letter(c)].width = w

out = ROOT / "output" / "私立高校_Web出願_進捗管理.xlsx"
wb.save(out)
print(out)
